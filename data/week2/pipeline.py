"""
pipeline.py
-----------
Week 2 — Data Engineering: Data Pipeline connecting ingestion -> text
extraction -> (pluggable) PII masking -> ready-for-storage clean text.

Cross-team dependency note (see roadmap, Week 2 table):
    "DevOps (PII Masking Engine) blocks Data Eng — masking must be callable
    before the ingestion pipeline stores data."

DevOps's masking service isn't built yet as of this pipeline being written,
so this module defines a `MaskingService` protocol (an interface) that the
real masking function can be dropped into once it exists — nothing here is
faked or marked "done" for masking itself. Everything up to and including
"call the masking function" is fully implemented and tested; only the
masking implementation is a stand-in.
"""

from dataclasses import dataclass
from typing import Protocol

try:
    from .text_extraction import extract_text, ExtractedDocument, TextExtractionError
except ImportError:
    from text_extraction import extract_text, ExtractedDocument, TextExtractionError


class MaskingService(Protocol):
    """
    Interface DevOps's real PII Masking Engine must satisfy to plug into
    this pipeline. Takes raw text, returns text with PII replaced by
    placeholders (e.g. "[NAME]", "[SSN]", "[EMAIL]").
    """
    def mask(self, text: str) -> str:
        ...


class PassthroughMaskingService:
    """
    Placeholder implementation used until DevOps's real masking service is
    ready. Does NOT mask anything — it exists only so the rest of the
    pipeline (extraction -> masking -> clean text) can be built, tested, and
    wired together now, without blocking on DevOps.

    IMPORTANT: This must be swapped for the real masking service before any
    real document text reaches storage or an external AI API — per the
    roadmap's Security & Operations requirement ("No unmasked PII ever
    reaches the third-party API").
    """
    def mask(self, text: str) -> str:
        return text


@dataclass
class PipelineResult:
    file_type: str
    masked_text: str
    tables: list
    page_or_sheet_count: int


def run_pipeline(file_type: str, file_bytes: bytes, masking_service: MaskingService) -> PipelineResult:
    """
    Full Week 2 flow: extract text from the raw file, then pass it through
    the given masking service before it's considered ready for storage.

    `masking_service` is injected (not hardcoded) so tests and early
    development can use `PassthroughMaskingService`, while production wiring
    swaps in DevOps's real implementation without changing this function.
    """
    try:
        extracted: ExtractedDocument = extract_text(file_type, file_bytes)
    except TextExtractionError as e:
        raise TextExtractionError(f"Pipeline failed at extraction step: {e}") from e

    masked_text = masking_service.mask(extracted.raw_text)

    return PipelineResult(
        file_type=extracted.file_type,
        masked_text=masked_text,
        tables=extracted.tables,
        page_or_sheet_count=extracted.page_or_sheet_count,
    )


# ---------------------------------------------------------------------------
# Quick manual test (run: python pipeline.py)
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    import io
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_paragraph("Client name: Jane Doe. SSN: 123-45-6789.")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    result = run_pipeline("docx", docx_bytes, masking_service=PassthroughMaskingService())

    print("Pipeline result (using PassthroughMaskingService — no real masking yet):")
    print("File type:", result.file_type)
    print("Masked text (unmasked for now, pending DevOps):", result.masked_text)
    print("Page/sheet count:", result.page_or_sheet_count)
